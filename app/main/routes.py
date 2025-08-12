"""
Main Routes for JobMate
Landing page, dashboards, and core application routes
"""

from flask import render_template, redirect, url_for, request, jsonify, current_app, send_file, session, flash
from flask_login import current_user, login_required, login_user
from app.main import bp
from app import db
from app.models.user import User
from app.models.application import Application
from app.models.job_posting import JobPosting
from app.utils import split_answer_and_code
from datetime import datetime, timedelta
from bson import ObjectId
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from app.ai_agents.gemini_utils import call_gemini_api
from app.models.resume import Resume
from werkzeug.utils import secure_filename
import os
import PyPDF2
import docx
import io
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.units import inch
from reportlab.lib.styles import ParagraphStyle
import re


def get_mongo_db():
    """Get MongoDB database connection"""
    return current_app.mongo_db


# --- MongoDB tailored resume helpers ---
def save_best_tailored_resume(mongo_db, user_id, job_id, ats_score, resume_text, cover_letter):
    collection = mongo_db.tailored_resumes
    existing = collection.find_one({"user_id": user_id, "job_id": job_id})
    # Ensure ats_score is treated as a number for comparison
    try:
        ats_score_num = float(ats_score) if ats_score is not None else 0
        existing_score = float(existing.get("ats_score", 0)) if existing else 0
        should_update = not existing or (ats_score_num > existing_score)
    except (ValueError, TypeError):
        # If conversion fails, always update
        should_update = True
        
    if should_update:
        collection.update_one(
            {"user_id": user_id, "job_id": job_id},
            {"$set": {
                "ats_score": ats_score,
                "resume_text": resume_text,
                "cover_letter": cover_letter,
                "updated_at": datetime.utcnow(),
                "created_at": existing.get("created_at", datetime.utcnow()) if existing else datetime.utcnow()
            }},
            upsert=True
        )

def get_best_tailored_resumes(mongo_db, user_id):
    return list(mongo_db.tailored_resumes.find({"user_id": user_id}))


def build_questions_data(generated_items):
    """
    Transforma la salida (generated_items) en la lista de dicts que la plantilla espera.
    generated_items: iterable de objetos/dicts desde el LLM o base de datos.
    """
    questions_data = []
    for item in generated_items:
        # Ajusta según la estructura real de `item`
        raw_expected = item.get("expected", "") or item.get("answer", "")
        answer_text, snippet, lang = split_answer_and_code(raw_expected)

        q = {
            "text": item.get("text", "") or item.get("question", ""),
            "relevance": item.get("relevance", "") or item.get("explanation", ""),
            "expected": answer_text,
            "code_snippet": snippet,
            "code_lang": (lang or "python")
        }
        questions_data.append(q)
    return questions_data


@bp.route('/')
def index():
    """Landing page"""
    # Redirect logged-in users to their dashboard
    if current_user.is_authenticated:
        if current_user.is_applicant():
            return redirect(url_for('main.applicant_dashboard'))
        else:
            return redirect(url_for('main.recruiter_dashboard'))
    
    return render_template('main/landing.html', title='JobMate - AI-Powered Job Matching')


@bp.route('/dashboard')
@login_required
def dashboard():
    """Generic dashboard redirect"""
    if current_user.is_applicant():
        return redirect(url_for('main.applicant_dashboard'))
    else:
        return redirect(url_for('main.recruiter_dashboard'))


@bp.route('/applicant/dashboard')
@login_required
def applicant_dashboard():
    """Enhanced applicant dashboard with working profile completion"""
    if not current_user.is_applicant():
        return redirect(url_for('main.recruiter_dashboard'))
    
    try:
        # Calculate profile completion
        profile_data = calculate_profile_completion(current_user)
        
        # Get user's active resume
        active_resume = current_user.get_active_resume()
        
        # Get recent applications using our optimized functions (limit to 5 for dashboard)
        pg_recent = get_postgres_applications(current_user)[:5]
        mongo_recent = get_mongodb_applications(current_user)[:5]
        
        # Combine and sort recent applications
        combined_recent_applications = pg_recent + mongo_recent
        
        # Sort by date and limit to 5 for dashboard
        def get_sort_date(app):
            applied_date = app.get('applied_date')
            return applied_date if applied_date else datetime.min
        
        combined_recent_applications.sort(key=get_sort_date, reverse=True)
        recent_applications = combined_recent_applications[:5]
        
        # Get job recommendations using the match module
        from app.models.job_posting import JobPosting
        recommended_jobs = []
        
        # Get active jobs for recommendations
        active_jobs = JobPosting.get_active_jobs(limit=10)
        if active_jobs and active_resume:
            try:
                from app.match.routes import calculate_job_match_score
                job_matches = []
                
                for job in active_jobs:
                    match_score = calculate_job_match_score(job, current_user, active_resume)
                    if match_score > 30:  # Only show jobs with reasonable match
                        job_matches.append({
                            'job': job,
                            'match_score': match_score
                        })
                
                # Sort by match score and limit to top 5 for dashboard
                recommended_jobs = sorted(job_matches, key=lambda x: x['match_score'], reverse=True)[:5]
            except ImportError:
                # If match module is not available, continue without recommendations
                pass
        
        # Fetch MongoDB jobs (limit 6)
        mongo_jobs = []
        try:
            mongo_db = current_app.mongo_db
            mongo_jobs = list(mongo_db.jobs.find({}, {"_id": 1, "title": 1, "company": 1, "description": 1}).limit(6))
            # Convert ObjectId to string for template
            for job in mongo_jobs:
                job['_id'] = str(job['_id'])
        except:
            # If MongoDB is not available, continue without mongo jobs
            pass
        
        return render_template('dashboard/applicant.html',
                             title='Dashboard',
                             user=current_user,
                             completion_percentage=profile_data['percentage'],
                             profile_completion_items=profile_data['items'],
                             missing_profile_items=profile_data['missing_items'],
                             active_resume=active_resume,
                             recent_applications=recent_applications,
                             recommended_jobs=recommended_jobs,
                             mongo_jobs=mongo_jobs)
        
    except Exception as e:
        current_app.logger.error(f"Dashboard error: {str(e)}")
        
        # Provide safe fallbacks
        return render_template('dashboard/applicant.html',
                             title='Dashboard',
                             user=current_user,
                             completion_percentage=0,
                             profile_completion_items=[],
                             missing_profile_items=[],
                             active_resume=None,
                             recent_applications=[],
                             recommended_jobs=[],
                             mongo_jobs=[])


@bp.route('/recruiter/dashboard')
@login_required
def recruiter_dashboard():
    """Recruiter dashboard"""
    if not current_user.is_recruiter():
        return redirect(url_for('main.applicant_dashboard'))
    
    # Get recruiter's job postings
    job_postings = current_user.job_postings.order_by(
        JobPosting.created_at.desc()
    ).limit(10).all()
    
    # Calculate statistics
    total_jobs = current_user.job_postings.count()
    active_jobs = current_user.job_postings.filter_by(status='active').count()
    
    # Get recent applications to recruiter's jobs
    job_ids = [job.id for job in current_user.job_postings.all()]
    recent_applications = []
    total_applications = 0
    new_applications = 0
    
    # Add application counts to each job posting
    seven_days_ago = datetime.utcnow() - timedelta(days=7)
    twenty_four_hours_ago = datetime.utcnow() - timedelta(hours=24)
    
    for job in job_postings:
        # Count total applications for this job
        job.total_applications = Application.query.filter_by(job_posting_id=job.id).count()
        
        # Count new applications (last 7 days) for this job
        job.new_applications = Application.query.filter(
            Application.job_posting_id == job.id,
            Application.created_at >= seven_days_ago
        ).count()
        
        # Calculate days since posted
        if job.created_at:
            job.days_ago = (datetime.utcnow() - job.created_at).days
        else:
            job.days_ago = 0
    
    if job_ids:
        recent_applications = Application.query.filter(
            Application.job_posting_id.in_(job_ids)
        ).order_by(Application.created_at.desc()).limit(5).all()
        
        total_applications = Application.query.filter(
            Application.job_posting_id.in_(job_ids),
            Application.status != 'rejected'
        ).count()
        
        # Applications from last 7 days
        new_applications = Application.query.filter(
            Application.job_posting_id.in_(job_ids),
            Application.created_at >= seven_days_ago
        ).count()
        
        # Applications from last 24 hours
        applications_24h = Application.query.filter(
            Application.job_posting_id.in_(job_ids),
            Application.created_at >= twenty_four_hours_ago
        ).count()
    else:
        applications_24h = 0
    
    # Mock data for demonstration
    stats = {
        'active_jobs': active_jobs,
        'total_applications': total_applications,
        'new_applications': new_applications,
        'applications_24h': applications_24h,
        'total_views': 1250,  # Mock data
        'interview_scheduled': 8  # Mock data
    }
    return render_template('dashboard/recruiter.html',
                         title='Recruiter Dashboard',
                         user=current_user,
                         job_postings=job_postings,
                         recent_applications=recent_applications,
                         current_date=datetime.now().strftime('%b %d'),
                         **stats)


@bp.route('/profile')
@login_required
def profile():
    """Redirect to enhanced profile page"""
    return redirect(url_for('user_profile.enhanced_profile'))

@bp.route('/profile/update', methods=['POST'])
@login_required
def update_profile():
    """Handle profile update form submission with detailed logging"""
    
    try:
        # Log form data received
        current_app.logger.info(f"=== Profile Update for User {current_user.id} ===")
        current_app.logger.info(f"Form data received: {dict(request.form)}")
        
        # Get form data
        first_name = request.form.get('first_name', '').strip()
        last_name = request.form.get('last_name', '').strip()
        phone = request.form.get('phone', '').strip()
        city = request.form.get('city', '').strip()
        bio = request.form.get('bio', '').strip()
        experience_level = request.form.get('experience_level', '').strip()
        skills = request.form.get('skills', '').strip()
        
        # Log what we're about to save
        current_app.logger.info("Data to be saved:")
        current_app.logger.info(f"  first_name: '{first_name}' (length: {len(first_name)})")
        current_app.logger.info(f"  last_name: '{last_name}' (length: {len(last_name)})")
        current_app.logger.info(f"  phone: '{phone}' (length: {len(phone)})")
        current_app.logger.info(f"  city: '{city}' (length: {len(city)})")
        current_app.logger.info(f"  bio: '{bio}' (length: {len(bio)})")
        current_app.logger.info(f"  experience_level: '{experience_level}' (length: {len(experience_level)})")
        current_app.logger.info(f"  skills: '{skills}' (length: {len(skills)})")
        
        # Update user fields
        current_user.first_name = first_name if first_name else None
        current_user.last_name = last_name if last_name else None
        current_user.phone = phone if phone else None
        current_user.city = city if city else None
        current_user.bio = bio if bio else None
        current_user.experience_level = experience_level if experience_level else None
        current_user.skills = skills if skills else None
        
        # Save to database
        db.session.commit()
        
        # Log what was actually saved
        current_app.logger.info("After database commit:")
        current_app.logger.info(f"  current_user.first_name: '{current_user.first_name}'")
        current_app.logger.info(f"  current_user.last_name: '{current_user.last_name}'")
        current_app.logger.info(f"  current_user.phone: '{current_user.phone}'")
        current_app.logger.info(f"  current_user.city: '{current_user.city}'")
        current_app.logger.info(f"  current_user.bio: '{current_user.bio}'")
        current_app.logger.info(f"  current_user.experience_level: '{current_user.experience_level}'")
        current_app.logger.info(f"  current_user.skills: '{current_user.skills}'")
        
        # Store timestamp for debugging
        session['last_profile_update'] = datetime.now().isoformat()
        
        # Calculate completion after save
        profile_data = calculate_profile_completion(current_user)
        current_app.logger.info(f"Profile completion after save: {profile_data['percentage']}%")
        
        flash('Profile updated successfully!', 'success')
        return redirect(url_for('main.applicant_dashboard'))
        
    except Exception as e:
        current_app.logger.error(f"Profile update error: {str(e)}", exc_info=True)
        db.session.rollback()
        flash('Error updating profile. Please try again.', 'error')
        return redirect(url_for('main.profile'))


def normalize_application_date(date_value):
    """Standardize date formats across data sources"""
    if date_value is None:
        return None
    
    if isinstance(date_value, str):
        try:
            # Handle various ISO format variations
            date_str = date_value.replace('Z', '+00:00')
            return datetime.fromisoformat(date_str)
        except (ValueError, TypeError):
            current_app.logger.warning(f"Failed to parse date string: {date_value}")
            return None
    
    # Already a datetime object
    if hasattr(date_value, 'year'):
        return date_value
    
    return None


def get_postgres_applications(user):
    """Fetch and format PostgreSQL applications with error handling"""
    applications = []
    
    try:
        pg_applications = user.applications.order_by(
            Application.created_at.desc()
        ).all()
        
        current_app.logger.info(f"Found {len(pg_applications)} PostgreSQL applications for user {user.id}")
        
        for app in pg_applications:
            # Get job title and company with fallback hierarchy
            job_title = app.job_title
            company_name = app.company_name
            
            if not job_title and app.job_posting:
                job_title = app.job_posting.title
            if not company_name and app.job_posting:
                company_name = app.job_posting.company_name
                
            # Fallback to defaults if still None
            job_title = job_title or 'Job Title Not Available'
            company_name = company_name or 'Company Not Available'
            
            applications.append({
                'id': f"pg_{app.id}",  # Prefix to avoid ID conflicts
                'job_title': job_title,
                'company_name': company_name,
                'applied_date': normalize_application_date(app.applied_at or app.created_at),
                'status': app.status or 'applied',
                'application_type': 'recruiter',
                'job_posting_id': app.job_posting_id,
                'ats_score': getattr(app, 'ats_score', None),
                'source': 'Internal'
            })
            
    except Exception as e:
        current_app.logger.error(f"PostgreSQL applications fetch failed: {str(e)}")
        # Don't re-raise, just log and continue with empty list
        
    return applications


def get_mongodb_applications(user):
    """Fetch and format MongoDB applications with optimized queries"""
    applications = []
    
    try:
        mongo_db = current_app.mongo_db
        if mongo_db is None:
            current_app.logger.warning("MongoDB connection not available")
            return applications
            
        # Single query with OR condition instead of multiple queries
        mongo_applications = list(
            mongo_db.job_applications.find({
                '$or': [
                    {'user_id': str(user.id)},
                    {'user_id': user.id}
                ]
            }).sort('applied_at', -1)
        )
        
        current_app.logger.info(f"Found {len(mongo_applications)} MongoDB applications for user {user.id}")
        
        # Batch fetch job details to reduce database calls
        job_ids = [app.get('job_id') for app in mongo_applications if app.get('job_id')]
        job_details_map = {}
        
        if job_ids:
            try:
                from bson import ObjectId
                # Convert string IDs to ObjectIds for batch query
                object_ids = []
                for job_id in job_ids:
                    try:
                        if not str(job_id).startswith('recruiter_'):
                            object_ids.append(ObjectId(job_id))
                    except:
                        continue
                
                if object_ids:
                    job_details = mongo_db.jobs.find({'_id': {'$in': object_ids}})
                    job_details_map = {str(job['_id']): job for job in job_details}
                    
            except Exception as e:
                current_app.logger.warning(f"Failed to batch fetch job details: {str(e)}")
        
        # Remove duplicates by _id (in case of data inconsistency)
        seen_ids = set()
        
        for app in mongo_applications:
            app_id = str(app.get('_id'))
            if app_id in seen_ids:
                continue
            seen_ids.add(app_id)
            
            # Get job details from our batch-fetched map
            job_id = app.get('job_id')
            job_details = job_details_map.get(str(job_id)) if job_id else None
            
            # Build application data with fallbacks
            job_title = (
                app.get('job_title') or 
                (job_details.get('title') if job_details else None) or 
                'External Job'
            )
            
            company_name = (
                app.get('company_name') or 
                (job_details.get('company') if job_details else None) or 
                'External Company'
            )
            
            applications.append({
                'id': f"mg_{app_id}",  # Prefix to avoid ID conflicts
                'job_title': job_title,
                'company_name': company_name,
                'applied_date': normalize_application_date(app.get('applied_at')),
                'status': app.get('status', 'applied'),
                'application_type': 'external',
                'job_posting_id': job_id,
                'ats_score': app.get('ats_score'),
                'source': 'External'
            })
            
    except Exception as e:
        current_app.logger.error(f"MongoDB applications fetch failed: {str(e)}")
        # Don't re-raise, just log and continue with empty list
        
    return applications


@bp.route('/applications')
@login_required
def applications():
    """Enhanced user applications page with improved error handling and performance"""
    
    if not current_user.is_applicant():
        flash('Access denied. Applicant account required.', 'error')
        return redirect(url_for('main.dashboard'))
    
    try:
        # Fetch applications from both sources
        pg_applications = get_postgres_applications(current_user)
        mongo_applications = get_mongodb_applications(current_user)
        
        # Combine applications
        combined_applications = pg_applications + mongo_applications
        
        # Sort by date (newest first) with safe date handling
        def get_sort_date(app):
            applied_date = app.get('applied_date')
            return applied_date if applied_date else datetime.min
        
        combined_applications.sort(key=get_sort_date, reverse=True)
        
        current_app.logger.info(
            f"Applications page loaded: {len(pg_applications)} internal, "
            f"{len(mongo_applications)} external, {len(combined_applications)} total"
        )
        
        return render_template(
            'main/applications.html',
            title='My Applications',
            applications=combined_applications,
            total_count=len(combined_applications),
            internal_count=len(pg_applications),
            external_count=len(mongo_applications)
        )
        
    except Exception as e:
        current_app.logger.error(f"Critical error in applications route: {str(e)}")
        flash('Unable to load applications. Please try again later.', 'error')
        return redirect(url_for('main.dashboard'))


@bp.route('/about')
def about():
    """About page"""
    return render_template('main/about.html', title='About JobMate')


@bp.route('/features')
def features():
    """Features page"""
    return render_template('main/features.html', title='Features')


@bp.route('/pricing')
def pricing():
    """Pricing page"""
    return render_template('main/pricing.html', title='Pricing')


@bp.route('/contact')
def contact():
    """Contact page"""
    return render_template('main/contact.html', title='Contact Us')


@bp.route('/privacy')
def privacy():
    """Privacy policy page"""
    return render_template('main/privacy.html', title='Privacy Policy')


@bp.route('/terms')
def terms():
    """Terms of service page"""
    return render_template('main/terms.html', title='Terms of Service')


# API Routes
@bp.route('/api/dashboard/stats')
@login_required
def dashboard_stats():
    """Get dashboard statistics for current user"""
    
    if current_user.is_applicant():
        stats = {
            'total_applications': current_user.applications.count(),
            'pending_applications': current_user.applications.filter_by(status='pending').count(),
            'interviews_scheduled': current_user.applications.filter_by(status='interview').count(),
            'profile_completion': calculate_profile_completion(current_user),
            'active_resume': current_user.get_active_resume() is not None
        }
    else:  # recruiter
        stats = {
            'total_job_postings': current_user.job_postings.count(),
            'active_job_postings': current_user.job_postings.filter_by(status='active').count(),
            'total_applications': 0,  # TODO: Calculate applications to recruiter's jobs
            'new_applications': 0,    # TODO: Calculate new applications this week
            'candidates_reviewed': 0  # TODO: Calculate reviewed candidates
        }
    
    return jsonify(stats)


@bp.route('/api/available-jobs')
@login_required
def available_jobs():
    """Get list of available job postings for ATS analysis"""
    try:
        # Get active job postings with basic info
        jobs = JobPosting.query.filter_by(status='active').order_by(
            JobPosting.created_at.desc()
        ).limit(20).all()
        
        jobs_data = []
        for job in jobs:
            jobs_data.append({
                'id': job.id,
                'title': job.title,
                'company_name': job.company_name,
                'location': job.location,
                'description': job.description[:200] if job.description else '',
                'experience_level': job.experience_level,
                'job_type': job.job_type,
                'created_at': job.created_at.isoformat() if job.created_at else None
            })
        
        return jsonify({'jobs': jobs_data})
        
    except Exception as e:
        return jsonify({'error': str(e), 'jobs': []}), 500


@bp.route('/api/search')
def search():
    """Global search endpoint"""
    query = request.args.get('q', '').strip()
    search_type = request.args.get('type', 'all')  # jobs, companies, candidates
    
    if not query:
        return jsonify({'results': []})
    
    results = []
    
    if search_type in ['all', 'jobs']:
        # TODO: Search jobs in MongoDB
        pass
    
    if search_type in ['all', 'companies']:
        # TODO: Search companies
        pass
    
    if current_user.is_authenticated and current_user.is_recruiter() and search_type in ['all', 'candidates']:
        # TODO: Search candidates (only for recruiters)
        pass
    
    return jsonify({'results': results, 'total': len(results)})


@bp.route('/mongo-test')
def mongo_test():
    """Test MongoDB connection and show job count and some job details"""
    mongo_db = current_app.mongo_db
    count = mongo_db.jobs.count_documents({})
    # Fetch up to 5 job documents, only show _id and title if available
    jobs = list(mongo_db.jobs.find({}, {"_id": 1, "title": 1}).limit(5))
    job_list = "<ul>" + "".join([
        f"<li>{str(job.get('_id'))}: {job.get('title', '(no title)')}</li>" for job in jobs
    ]) + "</ul>" if jobs else "<p>No jobs found.</p>"
    return f"Job count in MongoDB: {count}<br>Sample jobs:{job_list}"


@bp.route('/mongo-jobs')
@login_required
def mongo_jobs():
    """Show jobs from MongoDB for tailoring resume/cover letter"""
    mongo_db = current_app.mongo_db
    jobs = list(mongo_db.jobs.find({}, {"_id": 1, "title": 1, "company": 1, "description": 1}))
    return render_template('mongo_jobs.html', jobs=jobs)


def extract_text_from_file(file_path, file_ext):
    if file_ext == '.pdf':
        text = ''
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() or ''
        return text
    elif file_ext == '.docx':
        doc = docx.Document(file_path)
        return '\n'.join([para.text for para in doc.paragraphs])
    elif file_ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    else:
        raise ValueError('Unsupported file type')


def postprocess_tailored_resume_output(ai_result, user_resume_sections=None):
    """
    Post-process and validate the AI output for tailored resume.
    - Ensures all required sections are present.
    - Fixes/corrects LinkedIn/contact fields.
    - Deduplicates and merges with user data if needed.
    - Cleans up formatting and section order.
    - Removes duplicate section headers.
    """
    import re
    from app.jobs.routes import remove_duplicate_section_headers
    
    required_sections = ["Summary", "Skills", "Experience", "Education"]
    resume_text = ai_result.get("tailored_resume", "")
    
    # First, remove any duplicate section headers
    resume_text = remove_duplicate_section_headers(resume_text)
    
    # Parse into sections
    def parse_sections(text):
        section_pattern = re.compile(r"^([A-Za-z &]+):$", re.MULTILINE)
        sections = {}
        last = None
        for line in text.splitlines():
            m = section_pattern.match(line.strip())
            if m:
                last = m.group(1).strip()
                sections[last] = []
            elif last:
                sections[last].append(line)
        return sections
    ai_sections = parse_sections(resume_text)
    # Merge with user resume if provided
    if user_resume_sections:
        for sec in required_sections:
            if sec not in ai_sections and sec.upper() in user_resume_sections:
                ai_sections[sec] = user_resume_sections[sec.upper()]
    # Ensure all required sections exist
    for sec in required_sections:
        if sec not in ai_sections:
            ai_sections[sec] = ["(Not provided)"]
    # Fix LinkedIn/contact fields in header
    header_lines = []
    for k in ["Name", "Contact", "Location"]:
        found = False
        for line in resume_text.splitlines():
            if line.strip().lower().startswith(f"{k.lower()}:"):
                val = line.split(":", 1)[-1].strip()
                if k == "Contact":
                    # Try to fix LinkedIn if missing protocol
                    val = re.sub(r"(linkedin\\.com/[^ |]+)", r"https://\\1", val)
                header_lines.append(f"{k}: {val}")
                found = True
                break
        if not found:
            header_lines.append(f"{k}: (Not provided)")
    # Rebuild resume text in canonical order
    section_order = ["Summary", "Skills", "Experience", "Education"] + [k for k in ai_sections if k not in required_sections]
    rebuilt = []
    rebuilt.append("---")
    rebuilt.extend(header_lines)
    rebuilt.append("")
    for sec in section_order:
        rebuilt.append(f"{sec}:")
        rebuilt.extend([l for l in ai_sections[sec] if l.strip()])
        rebuilt.append("")
    rebuilt.append("---")
    ai_result["tailored_resume"] = "\n".join(rebuilt)
    return ai_result


@bp.route('/tailor/<job_id>', methods=['GET', 'POST'])
@login_required
def tailor_resume(job_id):
    """Tailor resume or cover letter for a specific job using Gemini AI - supports both MongoDB and PostgreSQL jobs"""
    mongo_db = current_app.mongo_db
    job = None
    
    # Handle both recruiter jobs (PostgreSQL) and external jobs (MongoDB)
    try:
        if job_id.startswith('recruiter_'):
            # This is a PostgreSQL recruiter job
            from app.models.job_posting import JobPosting
            pg_job_id = job_id.replace('recruiter_', '')
            pg_job = JobPosting.query.filter_by(id=pg_job_id, status='active').first()
            
            if pg_job:
                # Convert PostgreSQL job to consistent format
                job = {
                    '_id': job_id,
                    'title': pg_job.title,
                    'company': pg_job.company_name,
                    'company_name': pg_job.company_name,
                    'location': pg_job.location,
                    'description': pg_job.description,
                    'requirements': pg_job.requirements,
                    'job_type': pg_job.employment_type,
                    'experience_level': pg_job.experience_level,
                    'salary_min': pg_job.salary_min,
                    'salary_max': pg_job.salary_max,
                    'remote_type': pg_job.remote_type,
                    'is_recruiter_job': True,
                    'job_posting_id': pg_job.id
                }
        else:
            # This is a MongoDB external job
            job = mongo_db.jobs.find_one({"_id": ObjectId(job_id)})
            if job:
                job['is_recruiter_job'] = False
                
    except Exception as e:
        current_app.logger.error(f"Error fetching job {job_id}: {str(e)}")
        job = None
    
    if not job:
        return "Job not found", 404

    tailored_result = None
    ats_score = None
    error = None
    resume_text = None
    
    # Get application URL from job data using the correct field name
    application_url = job.get('job_url_direct') or job.get('application_url') or job.get('company_website')
    application_email = job.get('email') or job.get('application_email') or job.get('contact_email')

    if request.method == 'POST':
        # Handle file upload
        if 'resume_file' not in request.files or request.files['resume_file'].filename == '':
            error = 'Please upload a resume file.'
        else:
            resume_file = request.files['resume_file']
            filename = secure_filename(resume_file.filename)
            file_ext = os.path.splitext(filename)[1].lower()
            if file_ext not in ['.pdf', '.docx', '.txt']:
                error = 'Unsupported file type. Please upload PDF, DOCX, or TXT.'
            else:
                temp_path = os.path.join('/tmp', filename)
                resume_file.save(temp_path)
                try:
                    resume_text = extract_text_from_file(temp_path, file_ext)
                    session['user_resume_text'] = resume_text  # Save user resume text for later PDF export
                except Exception as e:
                    error = f'Failed to extract text from resume: {e}'
                finally:
                    os.remove(temp_path)

        if resume_text and not error:
            # Prepare prompt for Gemini
            prompt = f"""
# RESUME TAILORING TASK

You are an expert ATS optimization specialist and resume writer with deep expertise in the tech sector. Your task is to tailor the candidate's resume to maximize ATS match score and appeal to hiring managers for the specific job description.

## JOB DESCRIPTION
```
{job.get('description', '')}
```

## ORIGINAL RESUME
```
{resume_text}
```

## TAILORING INSTRUCTIONS

1. MAINTAIN ACCURACY: Never fabricate experience, education, or skills not mentioned in the original resume.

2. KEYWORD OPTIMIZATION:
   - Identify primary and secondary keywords from the job description
   - Naturally integrate these keywords into appropriate sections of the resume
   - Match phrasing used in the job description where appropriate
   - Ensure keyword density is optimized but not excessive (aim for 5-8 exact matches per key requirement)
   - Focus especially on technical skills, tools, and methodologies mentioned in the job

3. FORMATTING & STRUCTURE:
   - Use clean, ATS-friendly formatting
   - Prioritize relevant experience and skills based on job requirements
   - Use standard section headings that ATS systems recognize
   - Maintain chronological order within sections
   - Use bullet points for readability
   - Format the resume using the specified structure below

4. CONTENT ENHANCEMENT:
   - Transform generic statements into achievement-focused bullets
   - For EACH job experience, include EXACTLY 5 bullet points - no more, no less
   - Each bullet point MUST include metrics (numbers, percentages, dollar amounts) and demonstrate clear impact
   - Use strong action verbs at the beginning of bullet points (e.g., Increased, Developed, Implemented)
   - Cover different aspects in the 5 bullets: technical skills, leadership, problem-solving, collaboration, and quantifiable business impact
   - Organize the Skills section with clear subheadings (Programming Languages, Frameworks & Libraries, etc.)
   - Under each skill subheading, list specific skills separated by commas (not as bullet points)
   - Match skills specifically to those mentioned in the job description, prioritizing exact keyword matches
   - Highlight transferable skills that match the job requirements
   - Emphasize Ontario tech sector experience if available

5. COVER LETTER REQUIREMENTS:
   - Create a professional business letter format with proper contact information layout
   - Place candidate's contact information (name, email, phone, address) at the top left
   - Include today's date below contact information
   - Add hiring manager and company information below the date
   - Use formal business letter structure with clear paragraphs
   - Reference specific achievements and metrics from the tailored resume
   - Keep it concise but impactful (3-4 paragraphs)
   - Include relevant keywords from the job description naturally

## REQUIRED OUTPUT FORMAT
Format the tailored resume as plain text in the following structure:

---
Name: <Full Name>
Contact: <Email> | <Phone> | <LinkedIn>
Location: <City, Province>

Summary:
<Professional summary paragraph highlighting key qualifications relevant to job>

Skills:
Programming Languages: <List key programming languages from resume that match job requirements>
Frameworks & Libraries: <List relevant frameworks and libraries that match job requirements>
Tools & Technologies: <List relevant tools, platforms, and technologies from resume>
Cloud & Infrastructure: <List cloud platforms (AWS, Azure, GCP) and infrastructure tools (Terraform, Kubernetes) that match job requirements>
Databases & Storage: <List databases (SQL, NoSQL, PostgreSQL, MongoDB) and data storage solutions from resume>
Domain Knowledge: <List industry-specific skills and knowledge areas>
Soft Skills: <List 3-5 most relevant soft skills for the position>

Experience:
<Company Name> | <Job Title> | <Start Date> - <End Date>
- <Quantified achievement with specific metrics (e.g., "Increased application performance by 40% through database optimization")>
- <Technical achievement demonstrating skills required in job posting with measurable outcome>
- <Leadership or collaborative achievement showing teamwork with quantifiable results>
- <Problem-solving achievement with clear business impact and metrics>
- <Innovation or project accomplishment directly relevant to job requirements with measurable success>

# REPEAT THE ABOVE FORMAT WITH EXACTLY 5 BULLET POINTS FOR EACH JOB EXPERIENCE

Education:
<Degree> | <Institution> | <Graduation Year>

---

## OUTPUT REQUIREMENTS

Provide your response in JSON format with the following structure:

```json
{{
  "tailored_resume": "The complete tailored resume text with appropriate formatting",
  "cover_letter": "A professional business format cover letter with proper contact structure - see format requirements below",
  "ats_score": "A number between 0-100 representing the estimated ATS match score",
  "suggestions": [
    {{
      "area": "Specific section or aspect of the resume",
      "suggestion": "Detailed actionable suggestion to further improve the resume"
    }}
    // Include 5-8 specific suggestions
  ]
}}
```

## COVER LETTER FORMAT REQUIREMENTS

The cover letter MUST follow this professional business letter format using ACTUAL information from the resume:

```
[Extract candidate's full name from resume]
[Extract candidate's email from resume]
[Extract candidate's phone from resume]  
[Extract candidate's address from resume]

[Current date in format: Month Day, Year]

{job.get('company_name', 'Hiring Company')}
{job.get('location', '')}

Dear Hiring Manager,

[Opening paragraph: Express interest in the specific position "{job.get('title', 'the position')}" at {job.get('company_name', 'your company')}]

[Body paragraph 1: Highlight 2-3 key qualifications from your resume that directly match the job requirements, using specific examples and metrics where possible]

[Body paragraph 2: Demonstrate knowledge of {job.get('company_name', 'the company')} and explain why you want to work there specifically, connecting your experience to their needs]

[Closing paragraph: Express enthusiasm for an interview, mention that your resume is attached, and provide a professional closing]

Sincerely,

[Extract and use candidate's full name from resume]
```

Important cover letter guidelines:
- Use the candidate's actual contact information from their resume
- Reference specific qualifications and achievements from the tailored resume
- Keep it concise (3-4 paragraphs maximum)
- Use keywords from the job description naturally
- Maintain a professional, confident tone
- Include specific metrics and achievements when possible

For the suggestions, focus on:
1. Areas where keywords could be better integrated
2. Achievements that could be better quantified
3. Skills that should be emphasized based on the job
4. Format improvements for ATS optimization
5. Section ordering or content prioritization

Be specific and actionable with each suggestion.
"""
            gemini_response = call_gemini_api(prompt)
            import json
            if 'candidates' in gemini_response:
                try:
                    content = gemini_response['candidates'][0]['content']['parts'][0]['text']
                    # Try to parse as JSON, but if it fails, attempt to extract JSON from text
                    try:
                        result = json.loads(content)
                        tailored_result = result
                        # Ensure ats_score is a number
                        ats_score_raw = result.get('ats_score')
                        if isinstance(ats_score_raw, str):
                            # Try to extract numeric value if it's a string like "85/100" or just "85"
                            ats_score_match = re.search(r'(\d+)', ats_score_raw)
                            ats_score = int(ats_score_match.group(1)) if ats_score_match else None
                        else:
                            # Keep as is if it's already a number
                            ats_score = ats_score_raw
                    except Exception as e:
                        import re
                        match = re.search(r'\{.*\}', content, re.DOTALL)
                        if match:
                            try:
                                result = json.loads(match.group(0))
                                tailored_result = result
                                # Ensure ats_score is a number
                                ats_score_raw = result.get('ats_score')
                                if isinstance(ats_score_raw, str):
                                    # Try to extract numeric value if it's a string like "85/100" or just "85"
                                    ats_score_match = re.search(r'(\d+)', ats_score_raw)
                                    ats_score = int(ats_score_match.group(1)) if ats_score_match else None
                                else:
                                    # Keep as is if it's already a number
                                    ats_score = ats_score_raw
                            except Exception as e2:
                                # JSON parsing failed completely - use fallback approach
                                print(f"JSON parsing failed: {e2}")
                                # Try to extract text content even if JSON is malformed
                                tailored_text = content
                                # Remove JSON formatting attempts
                                tailored_text = re.sub(r'^\s*\{.*?"tailored_resume"\s*:\s*"', '', tailored_text, flags=re.DOTALL)
                                tailored_text = re.sub(r'",?\s*"cover_letter".*?\}?\s*$', '', tailored_text, flags=re.DOTALL)
                                tailored_text = tailored_text.replace('\\"', '"').replace('\\n', '\n')
                                
                                if tailored_text and len(tailored_text.strip()) > 100:
                                    # Create a basic result structure
                                    tailored_result = {
                                        'tailored_resume': tailored_text.strip(),
                                        'cover_letter': 'Due to technical issues, please generate a cover letter separately.',
                                        'ats_score': None
                                    }
                                    # Calculate ATS score using our enhanced algorithm
                                    from app.jobs.routes import calculate_enhanced_ats_score
                                    try:
                                        ats_score = calculate_enhanced_ats_score(
                                            tailored_text, 
                                            job.get('description', '') or job.get('summary', '') or str(job)
                                        )
                                        tailored_result['ats_score'] = ats_score
                                    except:
                                        ats_score = 75  # Default reasonable score
                                        tailored_result['ats_score'] = ats_score
                                else:
                                    error = f"AI response parsing failed. Please try again."
                        else:
                            error = f"Invalid AI response format. Please try again."
                except Exception as e:
                    error = f"Failed to extract Gemini response: {e}"
            else:
                error = gemini_response.get('error', 'Unknown error from Gemini API')

            # --- Enhanced ATS Optimization ---
            # If the current ATS score is below 90%, apply enhanced optimization
            if tailored_result and ats_score is not None and ats_score < 90:
                current_app.logger.info(f"Current ATS score {ats_score}% below 90%, applying enhanced optimization")
                try:
                    from app.services.resume_tailor import EnhancedResumeTailor
                    
                    # Initialize enhanced tailor
                    enhanced_tailor = EnhancedResumeTailor(
                        gemini_api_key=current_app.config.get('GEMINI_API_KEY'),
                        gemini_model=current_app.config.get('GEMINI_MODEL', 'gemini-2.5-flash')
                    )
                    
                    # Use the already tailored resume as input for enhancement
                    base_resume = tailored_result.get('tailored_resume', resume_text)
                    job_desc = job.get('description', '') or job.get('summary', '') or str(job)
                    
                    # Apply enhanced optimization
                    enhancement_result = enhanced_tailor.tailor_resume_for_high_ats(
                        base_resume, 
                        job_desc, 
                        target_score=90
                    )
                    
                    if enhancement_result['success'] and enhancement_result['ats_score'] > ats_score:
                        # Use the enhanced result
                        tailored_result['tailored_resume'] = enhancement_result['tailored_resume']
                        tailored_result['ats_score'] = enhancement_result['ats_score']
                        ats_score = enhancement_result['ats_score']
                        
                        current_app.logger.info(f"Enhanced ATS score achieved: {ats_score}%")
                        
                        # Add enhancement details to the result
                        tailored_result['enhancement_details'] = {
                            'enhanced': True,
                            'iterations_used': enhancement_result['iterations_used'],
                            'keywords_matched': enhancement_result['keywords_matched'],
                            'total_keywords': enhancement_result['total_keywords'],
                            'optimization_history': enhancement_result.get('optimization_history', [])
                        }
                    else:
                        current_app.logger.warning(f"Enhancement failed or did not improve score")
                        
                except Exception as e:
                    current_app.logger.error(f"Enhanced optimization failed: {str(e)}")
                    # Continue with original result if enhancement fails

            # --- Post-process tailored resume output ---
            if tailored_result:
                user_sections = extract_resume_sections(resume_text) if resume_text else None
                # TEMPORARILY DISABLED: This function adds "(Not provided)" and affects resume display
                # tailored_result = postprocess_tailored_resume_output(tailored_result, user_sections)
                
                # Get updated ATS score and ensure it's numeric
                ats_score_raw = tailored_result.get('ats_score')
                if isinstance(ats_score_raw, str):
                    # Try to extract numeric value if it's a string like "85/100" or just "85"
                    ats_score_match = re.search(r'(\d+)', ats_score_raw)
                    ats_score = int(ats_score_match.group(1)) if ats_score_match else None
                else:
                    # Keep as is if it's already a number
                    ats_score = ats_score_raw

            # --- Save best tailored resume in MongoDB ---
            if tailored_result and ats_score is not None:
                save_best_tailored_resume(
                    mongo_db,
                    str(current_user.id),
                    str(job_id),
                    ats_score,
                    tailored_result.get('tailored_resume', ''),
                    tailored_result.get('cover_letter', '')
                )

    return render_template(
        'tailor.html', 
        job=job,
        tailored_result=tailored_result, 
        ats_score=ats_score, 
        error=error, 
        user_resume_text=resume_text if resume_text else None,
        application_url=application_url,
        application_email=application_email
    )


@bp.route('/download-tailored/<job_id>/<doc_type>', methods=['POST'])
@login_required
def download_tailored(job_id, doc_type):
    """Download tailored resume or cover letter as a PDF file with improved formatting"""
    from app.jobs.routes import remove_duplicate_section_headers
    
    tailored_resume = request.form.get('tailored_resume')
    cover_letter = request.form.get('cover_letter')
    if doc_type == 'resume':
        content = tailored_resume or ''
        # Apply duplicate removal to ensure clean content
        content = remove_duplicate_section_headers(content)
        filename = f'tailored_resume_{job_id}.pdf'
    elif doc_type == 'cover_letter':
        content = cover_letter or ''
        filename = f'cover_letter_{job_id}.pdf'
    else:
        return "Invalid document type", 400

    # PDF generation with formatting
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.75*inch, rightMargin=0.75*inch, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    story = []

    # Handle cover letter differently from resume
    if doc_type == 'cover_letter':
        # Cover Letter specific formatting
        cover_letter_style = ParagraphStyle('CoverLetter', parent=styles['Normal'], fontSize=12, leading=18, spaceAfter=8, alignment=TA_LEFT)
        contact_header_style = ParagraphStyle('ContactHeader', parent=styles['Normal'], fontSize=12, leading=14, spaceAfter=6, alignment=TA_LEFT)
        date_style = ParagraphStyle('Date', parent=styles['Normal'], fontSize=12, leading=14, spaceAfter=12, spaceBefore=12, alignment=TA_LEFT)
        signature_style = ParagraphStyle('Signature', parent=styles['Normal'], fontSize=12, leading=14, spaceAfter=6, spaceBefore=12, alignment=TA_LEFT)
        
        lines = [l for l in content.splitlines() if l.strip()]
        
        # Parse cover letter content
        candidate_name = ""
        candidate_email = ""
        candidate_phone = ""
        candidate_address = ""
        date_line = ""
        company_info = []
        body_paragraphs = []
        in_body = False
        
        for i, line in enumerate(lines):
            line_strip = line.strip()
            
            # Skip empty lines at the beginning
            if not line_strip and not in_body:
                continue
                
            # Detect candidate contact information (first few lines)
            if i < 6 and not in_body:
                if '@' in line_strip and not candidate_email:
                    candidate_email = line_strip
                elif line_strip.replace('(', '').replace(')', '').replace('-', '').replace(' ', '').replace('+1', '').isdigit() and not candidate_phone:
                    candidate_phone = line_strip
                elif any(word in line_strip.lower() for word in ['street', 'avenue', 'road', 'drive', 'mississauga', 'ontario', 'on', 'canada']) and not candidate_address:
                    candidate_address = line_strip
                elif not candidate_name and len(line_strip.split()) <= 4:
                    # Likely the candidate's name
                    words = line_strip.split()
                    if len(words) >= 2 and all(word.replace('.', '').replace('-', '').isalpha() for word in words):
                        candidate_name = line_strip
            
            # Detect date (common formats)
            elif any(month in line_strip for month in ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']) or ('2023' in line_strip or '2024' in line_strip or '2025' in line_strip):
                date_line = line_strip
                
            # Detect company information (after date, before "Dear")
            elif date_line and not line_strip.lower().startswith('dear') and not in_body:
                if line_strip and 'sincerely' not in line_strip.lower():
                    company_info.append(line_strip)
            
            # Also collect company info if no date was found yet but we're before "Dear"
            elif not in_body and not line_strip.lower().startswith('dear') and not any(month in line_strip for month in ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']):
                # This might be company information if it's not contact info
                if line_strip and 'sincerely' not in line_strip.lower() and len(line_strip) > 3:
                    # Check if it's not candidate info and not a placeholder
                    if not ('@' in line_strip or 
                           line_strip.replace('(', '').replace(')', '').replace('-', '').replace(' ', '').replace('+1', '').isdigit() or
                           any(word in line_strip.lower() for word in ['street', 'avenue', 'road', 'drive', 'mississauga', 'ontario', 'on', 'canada']) or
                           line_strip.startswith('[') and line_strip.endswith(']')):  # Skip placeholders
                        company_info.append(line_strip)
                    
            # Start of letter body
            elif line_strip.lower().startswith('dear'):
                in_body = True
                body_paragraphs.append(line_strip)
                
            # Letter body content
            elif in_body:
                if line_strip.lower() == 'sincerely,':
                    body_paragraphs.append(line_strip)
                    # Don't automatically add name - let AI handle it
                elif line_strip:
                    body_paragraphs.append(line_strip)
        
        # Build the PDF story for cover letter
        # Candidate contact information (left-aligned at top)
        if candidate_name:
            story.append(Paragraph(f'<b>{candidate_name}</b>', contact_header_style))
        if candidate_email:
            story.append(Paragraph(candidate_email, contact_header_style))
        if candidate_phone:
            story.append(Paragraph(candidate_phone, contact_header_style))
        if candidate_address:
            story.append(Paragraph(candidate_address, contact_header_style))
            
        # Date
        if date_line:
            story.append(Paragraph(date_line, date_style))
        else:
            # Add current date if not found
            from datetime import datetime
            current_date = datetime.now().strftime("%B %d, %Y")
            story.append(Paragraph(current_date, date_style))
            
        # Company information
        for company_line in company_info:
            story.append(Paragraph(company_line, contact_header_style))
            
        if company_info:
            story.append(Spacer(1, 12))
            
        # Letter body
        for paragraph in body_paragraphs:
            if paragraph.strip():
                story.append(Paragraph(paragraph, cover_letter_style))
            else:
                story.append(Spacer(1, 6))
        
        # Complete the cover letter PDF generation
        story.append(Spacer(1, 16))
        doc.build(story)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/pdf',
            as_attachment=True,
            download_name=filename
        )
                
    else:
        # Resume formatting (existing code)
        # Custom styles
        name_style = ParagraphStyle('Name', parent=styles['Heading1'], fontSize=22, alignment=TA_CENTER, spaceAfter=10, spaceBefore=10)
        contact_style = ParagraphStyle('Contact', parent=styles['Normal'], fontSize=11, alignment=TA_CENTER, textColor=colors.black, spaceAfter=10)
        section_style = ParagraphStyle('Section', parent=styles['Heading2'], fontSize=13, textColor=colors.HexColor('#1a4a7c'), spaceBefore=18, spaceAfter=6, leading=16, fontName='Helvetica-Bold')
        normal_style = ParagraphStyle('Normal', parent=styles['Normal'], fontSize=11, leading=15, spaceAfter=6)
        bullet_style = ParagraphStyle('Bullet', parent=styles['Normal'], fontSize=11, leftIndent=15, bulletIndent=5, leading=15)

        # --- Render AI-optimized resume or cover letter only ---
        lines = [l for l in content.splitlines() if l.strip()]

        # --- Remove any '---' separator lines and extract name/contact/location ---
        name = ''
        contact = ''
        location = ''
        extra_contact = ''
        rest_lines = []
        header_processed = False
        
        for i, l in enumerate(lines):
            l_strip = l.strip()
            l_lower = l_strip.lower()
            
            if l_strip == '---':
                continue  # skip separator
                
            # Try to detect name, contact, and location more flexibly
            if l_lower.startswith('name:'):
                name = l_strip.split(':', 1)[-1].strip()
                header_processed = True
            elif l_lower.startswith('contact:'):
                contact = l_strip.split(':', 1)[-1].strip()
                header_processed = True
            elif l_lower.startswith('location:'):
                location = l_strip.split(':', 1)[-1].strip()
                header_processed = True
            elif l_lower.startswith('www') or l_lower.startswith('http'):
                if not contact:  # Only use as contact if we don't have one
                    extra_contact = l_strip
                else:
                    rest_lines.append(l_strip)
                header_processed = True
            elif not header_processed and i < 10:  # Check first 10 lines for header info
                # Smart detection for header information
                if '@' in l_strip and ('gmail' in l_lower or 'email' in l_lower or '.' in l_strip):
                    # This looks like an email/contact line
                    contact = l_strip
                elif any(word in l_lower for word in ['street', 'avenue', 'road', 'city', 'on', 'ontario', 'canada']) and len(l_strip.split()) <= 10:
                    # This looks like a location
                    location = l_strip
                elif not name and len(l_strip.split()) <= 4 and not any(word in l_lower for word in ['summary', 'skills', 'experience', 'education', 'professional']):
                    # This might be a name (short, not a section header)
                    # Additional check: if it contains typical name patterns
                    words = l_strip.split()
                    if len(words) >= 2 and all(word.replace('.', '').replace('-', '').isalpha() for word in words):
                        name = l_strip
                    else:
                        rest_lines.append(l_strip)
                else:
                    rest_lines.append(l_strip)
            else:
                rest_lines.append(l_strip)

        # --- Render name bold and centered at the top (no header) ---
        # Clean up any "(Not provided)" text
        if name and name.strip() != "(Not provided)":
            story.append(Paragraph(f'<b>{name}</b>', name_style))
        # Render contact/location line centered below name (no headers)
        # Clean contact and location data
        clean_contact = contact if contact and contact.strip() != "(Not provided)" else ""
        clean_location = location if location and location.strip() != "(Not provided)" else ""
        
        contact_line = ' | '.join(filter(None, [clean_contact, clean_location, extra_contact]))
        if contact_line:
            story.append(Paragraph(contact_line, contact_style))
        story.append(Spacer(1, 8))

        # --- Section parsing with experience subheader logic, grouped skills, and section reordering ---
        section_keywords = [
            'summary', 'professional summary', 'skills', 'projects', 'relevant projects',
            'experience', 'professional experience', 'education', 'community & interests', 'interests', 'project experience'
        ]
        def is_section_header(line):
            l = line.strip()
            if l.endswith(':'):
                return True
            if l.isupper() and len(l) > 3:
                return True
            l_clean = l.lower().rstrip(':').strip()
            return l_clean in section_keywords

        subtitle_style = ParagraphStyle('Subtitle', parent=styles['Normal'], fontSize=11, leading=15, spaceBefore=8, spaceAfter=4, fontName='Helvetica-Bold')

        # --- Parse all sections into a dict for reordering ---
        sections = {}
        current_section = None
        section_buffer = []
        for line in rest_lines:
            l_strip = line.strip()
            if is_section_header(l_strip):
                if current_section:
                    sections[current_section] = section_buffer
                    section_buffer = []
                current_section = l_strip
            elif current_section:
                section_buffer.append(line)
            else:
                # If not in a section, treat as intro
                sections.setdefault('intro', []).append(line)
        if current_section:
            sections[current_section] = section_buffer

        # --- Section order: SUMMARY, EXPERIENCE, then others as found ---
        ordered_section_titles = []
        for key in sections.keys():
            if key.lower().startswith('summary'):
                ordered_section_titles.append(key)
        for key in sections.keys():
            if key.lower().startswith('experience'):
                ordered_section_titles.append(key)
        for key in sections.keys():
            # Exclude SUMMARY, EXPERIENCE, and LANGUAGE sections
            if key.lower().startswith('summary') or key.lower().startswith('experience') or key.lower().startswith('language'):
                continue
            ordered_section_titles.append(key)

        def flush_section(title, buffer):
            if not buffer:
                return
            story.append(Spacer(1, 10))
            story.append(Paragraph(f'<b>{title.rstrip(":").upper()}</b>', section_style))
            # SKILLS section: handle subheadings and skill lists
            if title.lower().startswith('skills'):
                current_category = None
                for line in buffer:
                    line_strip = line.strip()
                    if not line_strip:
                        continue
                    
                    # Check if this is a skill category heading (ends with a colon)
                    if ':' in line_strip and not line_strip.startswith('-') and not line_strip.startswith('•'):
                        # This is a category heading like "Programming Languages: Python, Java"
                        parts = line_strip.split(':', 1)
                        if len(parts) == 2:
                            category = parts[0].strip()
                            skills = parts[1].strip()
                            
                            # Add the category as a subheading
                            story.append(Paragraph(f'<b>{category}</b>', subtitle_style))
                            
                            # Add the skills as normal text (not bullets)
                            if skills:
                                story.append(Paragraph(skills, normal_style))
                        else:
                            # Just a heading without skills yet
                            story.append(Paragraph(f'<b>{line_strip}</b>', subtitle_style))
                    elif line_strip.startswith('-') or line_strip.startswith('•'):
                        # Handle any bullet points that might still exist in the skills section
                        story.append(Paragraph(line_strip.lstrip('-•').strip(), bullet_style, bulletText='•'))
                    else:
                        # Plain text that isn't a category heading
                        story.append(Paragraph(line_strip, normal_style))
            # EXPERIENCE section: bold subheader for lines with ' | '
            elif title.lower().startswith('experience') or title.lower().startswith('professional experience'):
                for line in buffer:
                    line_strip = line.strip()
                    if '|' in line_strip and not line_strip.startswith('-'):
                        story.append(Paragraph(f'<b>{line_strip}</b>', normal_style))
                    elif line_strip.startswith('-') or line_strip.startswith('•'):
                        story.append(Paragraph(line_strip.lstrip('-•').strip(), bullet_style, bulletText='•'))
                    else:
                        story.append(Paragraph(line_strip, normal_style))
            # EDUCATION section: bold institution name
            elif title.lower().startswith('education'):
                for line in buffer:
                    line_strip = line.strip()
                    # Try to bold the institution name (after first '|')
                    if '|' in line_strip:
                        parts = [p.strip() for p in line_strip.split('|')]
                        if len(parts) >= 2:
                            # Assume institution is the second part
                            bolded = f"{parts[0]} | <b>{parts[1]}</b>"
                            if len(parts) > 2:
                                bolded += ' | ' + ' | '.join(parts[2:])
                            story.append(Paragraph(bolded, normal_style))
                        else:
                            story.append(Paragraph(line_strip, normal_style))
                    else:
                        story.append(Paragraph(line_strip, normal_style))
            # PROJECT EXPERIENCE section: bold project titles (first non-bullet line before bullets)
            elif title.lower().startswith('project experience') or title.lower().startswith('projects'):
                in_project = False
                for idx, line in enumerate(buffer):
                    line_strip = line.strip()
                    if not line_strip:
                        continue
                    if not (line_strip.startswith('-') or line_strip.startswith('•')):
                        # Bold project title
                        story.append(Paragraph(f'<b>{line_strip}</b>', normal_style))
                        in_project = True
                    else:
                        story.append(Paragraph(line_strip.lstrip('-•').strip(), bullet_style, bulletText='•'))
            else:
                for line in buffer:
                    line_strip = line.strip()
                    if line_strip.startswith('-') or line_strip.startswith('•'):
                        story.append(Paragraph(line_strip.lstrip('-•').strip(), bullet_style, bulletText='•'))
                    else:
                        story.append(Paragraph(line_strip, normal_style))

        # Render sections in the new order
        for section_title in ordered_section_titles:
            flush_section(section_title, sections[section_title])
        # Add extra space at end
        story.append(Spacer(1, 16))
        doc.build(story)
        buffer.seek(0)
        return send_file(
        buffer,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=filename
    )


@bp.route('/download-my-tailored/<job_id>/<doc_type>', methods=['GET'])
@login_required
def download_my_tailored(job_id, doc_type):
    """Download tailored resume or cover letter for a job from the resume section (MongoDB source)"""
    from app.jobs.routes import remove_duplicate_section_headers
    
    mongo_db = current_app.mongo_db
    tailored = mongo_db.tailored_resumes.find_one({"user_id": str(current_user.id), "job_id": str(job_id)})
    if not tailored:
        return "No tailored resume found for this job.", 404
    if doc_type == 'resume':
        content = tailored.get('resume_text', '')
        # Apply duplicate removal to ensure clean content
        content = remove_duplicate_section_headers(content)
        filename = f'tailored_resume_{job_id}.pdf'
    elif doc_type == 'cover_letter':
        content = tailored.get('cover_letter', '')
        filename = f'cover_letter_{job_id}.pdf'
    else:
        return "Invalid document type", 400
    # PDF generation (reuse logic from download_tailored)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.75*inch, rightMargin=0.75*inch, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    story = []
    name_style = ParagraphStyle('Name', parent=styles['Heading1'], fontSize=22, alignment=TA_CENTER, spaceAfter=10, spaceBefore=10)
    contact_style = ParagraphStyle('Contact', parent=styles['Normal'], fontSize=11, alignment=TA_CENTER, textColor=colors.black, spaceAfter=10)
    section_style = ParagraphStyle('Section', parent=styles['Heading2'], fontSize=13, textColor=colors.HexColor('#1a4a7c'), spaceBefore=18, spaceAfter=6, leading=16, fontName='Helvetica-Bold')
    normal_style = ParagraphStyle('Normal', parent=styles['Normal'], fontSize=11, leading=15, spaceAfter=6)
    bullet_style = ParagraphStyle('Bullet', parent=styles['Normal'], fontSize=11, leftIndent=15, bulletIndent=5, leading=15)
    lines = [l for l in content.splitlines() if l.strip()]
    name = ''
    contact = ''
    location = ''
    extra_contact = ''
    rest_lines = []
    for l in lines:
        l_strip = l.strip()
        l_lower = l_strip.lower()
        if l_strip == '---':
            continue  # skip separator
        if l_lower.startswith('name:'):
            name = l_strip.split(':', 1)[-1].strip()
        elif l_lower.startswith('contact:'):
            contact = l_strip.split(':', 1)[-1].strip()
        elif l_lower.startswith('location:'):
            location = l_strip.split(':', 1)[-1].strip()
        elif l_lower.startswith('www') or l_lower.startswith('http'):
            extra_contact = l_strip
        else:
            rest_lines.append(l_strip)
    if name:
        story.append(Paragraph(f'<b>{name}</b>', name_style))
    contact_line = ' | '.join(filter(None, [contact, location, extra_contact]))
    if contact_line:
        story.append(Paragraph(contact_line, contact_style))
    story.append(Spacer(1, 8))
    section_keywords = [
        'summary', 'professional summary', 'skills', 'projects', 'relevant projects',
        'experience', 'professional experience', 'education', 'community & interests', 'interests', 'project experience'
    ]
    def is_section_header(line):
        l = line.strip()
        if l.endswith(':'):
            return True
        if l.isupper() and len(l) > 3:
            return True
        l_clean = l.lower().rstrip(':').strip()
        return l_clean in section_keywords
    subtitle_style = ParagraphStyle('Subtitle', parent=styles['Normal'], fontSize=11, leading=15, spaceBefore=8, spaceAfter=4, fontName='Helvetica-Bold')
    sections = {}
    current_section = None
    section_buffer = []
    for line in rest_lines:
        l_strip = line.strip()
        if is_section_header(l_strip):
            if current_section:
                sections[current_section] = section_buffer
                section_buffer = []
            current_section = l_strip
        elif current_section:
            section_buffer.append(line)
        else:
            sections.setdefault('intro', []).append(line)
    if current_section:
        sections[current_section] = section_buffer
    ordered_section_titles = []
    for key in sections.keys():
        if key.lower().startswith('summary'):
            ordered_section_titles.append(key)
    for key in sections.keys():
        if key.lower().startswith('experience'):
            ordered_section_titles.append(key)
    for key in sections.keys():
        if key.lower().startswith('summary') or key.lower().startswith('experience') or key.lower().startswith('language'):
            continue
        ordered_section_titles.append(key)
    def flush_section(title, buffer):
        if not buffer:
            return
        story.append(Spacer(1, 10))
        story.append(Paragraph(f'<b>{title.rstrip(":").upper()}</b>', section_style))
        if title.lower().startswith('skills'):
            group = None
            group_buffer = []
            for line in buffer:
                line_strip = line.strip()
                if line_strip.startswith('-') and len(line_strip) > 2 and not line_strip[1].isspace():
                    if group and group_buffer:
                        story.append(Paragraph(f'<b>{group.lstrip("- ")}</b>', subtitle_style))
                        for skill in group_buffer:
                            story.append(Paragraph(skill.lstrip('-•').strip(), bullet_style, bulletText='•'))
                        group_buffer = []
                    group = line_strip
                elif line_strip.startswith('-') or line_strip.startswith('•'):
                    group_buffer.append(line_strip)
                elif line_strip:
                    group_buffer.append(line_strip)
            if group and group_buffer:
                story.append(Paragraph(f'<b>{group.lstrip("- ")}</b>', subtitle_style))
                for skill in group_buffer:
                    story.append(Paragraph(skill.lstrip('-•').strip(), bullet_style, bulletText='•'))
            elif group_buffer:
                for skill in group_buffer:
                    story.append(Paragraph(skill.lstrip('-•').strip(), bullet_style, bulletText='•'))
        elif title.lower().startswith('experience') or title.lower().startswith('professional experience'):
            for line in buffer:
                line_strip = line.strip()
                if '|' in line_strip and not line_strip.startswith('-'):
                    story.append(Paragraph(f'<b>{line_strip}</b>', normal_style))
                elif line_strip.startswith('-') or line_strip.startswith('•'):
                    story.append(Paragraph(line_strip.lstrip('-•').strip(), bullet_style, bulletText='•'))
                else:
                    story.append(Paragraph(line_strip, normal_style))
        elif title.lower().startswith('education'):
            for line in buffer:
                line_strip = line.strip()
                if '|' in line_strip:
                    parts = [p.strip() for p in line_strip.split('|')]
                    if len(parts) >= 2:
                        bolded = f"{parts[0]} | <b>{parts[1]}</b>"
                        if len(parts) > 2:
                            bolded += ' | ' + ' | '.join(parts[2:])
                        story.append(Paragraph(bolded, normal_style))
                    else:
                        story.append(Paragraph(line_strip, normal_style))
                else:
                    story.append(Paragraph(line_strip, normal_style))
        elif title.lower().startswith('project experience') or title.lower().startswith('projects'):
            in_project = False
            for idx, line in enumerate(buffer):
                line_strip = line.strip()
                if not line_strip:
                    continue
                if not (line_strip.startswith('-') or line_strip.startswith('•')):
                    story.append(Paragraph(f'<b>{line_strip}</b>', normal_style))
                    in_project = True
                else:
                    story.append(Paragraph(line_strip.lstrip('-•').strip(), bullet_style, bulletText='•'))
        else:
            for line in buffer:
                line_strip = line.strip()
                if line_strip.startswith('-') or line_strip.startswith('•'):
                    story.append(Paragraph(line_strip.lstrip('-•').strip(), bullet_style, bulletText='•'))
                else:
                    story.append(Paragraph(line_strip, normal_style))
    for section_title in ordered_section_titles:
        flush_section(section_title, sections[section_title])
    story.append(Spacer(1, 16))
    doc.build(story)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=filename
    )


def calculate_profile_completion(user):
    """Calculate profile completion percentage with detailed logging"""
    
    current_app.logger.info(f"=== Calculating Profile Completion for User {user.id} ===")
    
    completion_items = []
    total_possible = 0
    completed = 0
    
    # Essential profile fields to check
    profile_checks = [
        ('email', 'Email address', getattr(user, 'email', None)),
        ('first_name', 'First name', getattr(user, 'first_name', None)),
        ('last_name', 'Last name', getattr(user, 'last_name', None)),
        ('phone', 'Phone number', getattr(user, 'phone', None)),
        ('city', 'Location', getattr(user, 'city', None)),
        ('bio', 'Professional summary', getattr(user, 'bio', None)),
        ('skills', 'Skills', getattr(user, 'skills', None)),
        ('experience_level', 'Experience level', getattr(user, 'experience_level', None))
    ]
    
    # Check each field
    for field_name, display_name, field_value in profile_checks:
        total_possible += 1
        
        # Log the raw value
        current_app.logger.info(f"Checking {field_name}: raw value = '{field_value}' (type: {type(field_value)})")
        
        is_complete = False
        if field_value is not None:
            if isinstance(field_value, str):
                is_complete = field_value.strip() != ''
                current_app.logger.info(f"  String check: stripped = '{field_value.strip()}', is_complete = {is_complete}")
            elif isinstance(field_value, list):
                is_complete = len(field_value) > 0
                current_app.logger.info(f"  List check: length = {len(field_value)}, is_complete = {is_complete}")
            else:
                is_complete = True
                current_app.logger.info(f"  Other type check: is_complete = {is_complete}")
        else:
            current_app.logger.info(f"  Value is None: is_complete = {is_complete}")
        
        completion_items.append({
            'field': field_name,
            'name': display_name,
            'completed': is_complete,
            'value': field_value
        })
        
        if is_complete:
            completed += 1
            current_app.logger.info(f"  ✓ {display_name} completed")
        else:
            current_app.logger.info(f"  ✗ {display_name} missing")
    
    # Calculate percentage (resume upload not required for completion)
    completion_percentage = round((completed / total_possible) * 100) if total_possible > 0 else 0
    
    # Final logging
    current_app.logger.info(f"=== Profile Completion Summary ===")
    current_app.logger.info(f"Completed: {completed}/{total_possible} = {completion_percentage}%")
    for item in completion_items:
        status = "✓" if item['completed'] else "✗"
        current_app.logger.info(f"  {status} {item['name']}: '{item['value']}'")
    
    return {
        'percentage': completion_percentage,
        'completed': completed,
        'total': total_possible,
        'items': completion_items,
        'missing_items': [item for item in completion_items if not item['completed']]
    }

def extract_resume_sections(resume_text):
    """
    Parse resume text into a dict of sections: {section_name: [lines]}
    """
    section_headers = [
        'summary', 'professional summary', 'skills', 'projects', 'relevant projects',
        'experience', 'professional experience', 'work experience', 'education',
        'community & interests', 'interests', 'project experience', 'certifications', 'awards'
    ]
    def is_section_header(line):
        l = line.strip()
        if l.endswith(':'):
            return True
        if l.isupper() and len(l) > 3:
            return True
        l_clean = l.lower().rstrip(':').strip()
        return l_clean in section_headers
    sections = {}
    current_section = None
    section_buffer = []
    for line in resume_text.splitlines():
        l_strip = line.strip()
        if is_section_header(l_strip):
            if current_section:
                sections[current_section] = section_buffer
                section_buffer = []
            current_section = l_strip.rstrip(':').upper()
        elif current_section:
            section_buffer.append(line)
        else:
            sections.setdefault('INTRO', []).append(line)
    if current_section:
        sections[current_section] = section_buffer
    return sections

def extract_job_keywords(job_description):
    """
    Extract keywords and requirements from job description.
    Returns a dict with keys: skills, qualifications, responsibilities, keywords
    """
    # Simple regex/keyword-based extraction (can be improved with AI)
    skills = re.findall(r"[Ss]kills?[:\-\s]*([\w, .]+)", job_description)
    qualifications = re.findall(r"[Qq]ualifications?[:\-\s]*([\w, .]+)", job_description)
    responsibilities = re.findall(r"[Rr]esponsabilit(?:y|ies)[:\-\s]*([\w, .]+)", job_description)
    # Extract all unique words longer than 3 chars as keywords (can be improved)
    words = set(w.lower() for w in re.findall(r"\b\w{4,}\b", job_description))
    return {
        'skills': skills,
        'qualifications': qualifications,
        'responsibilities': responsibilities,
        'keywords': list(words)
    }


# @bp.route('/my-tailored-resumes')
# @login_required
# def my_tailored_resumes():
#     """Show all best tailored resumes and cover letters for the current user (from MongoDB)"""
#     if not current_user.is_applicant():
#         return redirect(url_for('main.dashboard'))
#     mongo_db = current_app.mongo_db
#     tailored_list = get_best_tailored_resumes(mongo_db, str(current_user.id))
#     # Fetch job titles for display
#     job_ids = [t['job_id'] for t in tailored_list]
#     jobs = {str(j['_id']): j for j in mongo_db.jobs.find({'_id': {'$in': [ObjectId(jid) for jid in job_ids]}})}
#     # Build display list
#     display_list = []
#     for t in tailored_list:
#         job = jobs.get(t['job_id'])
#         display_list.append({
#             'job_id': t['job_id'],
#             'job_title': job['title'] if job else '(Job not found)',
#             'ats_score': t.get('ats_score'),
#             'updated_at': t.get('updated_at'),
#             'resume_text': t.get('resume_text', ''),
#             'cover_letter': t.get('cover_letter', '')
#         })
#     return render_template('resume/my_resumes.html', tailored_resumes=display_list, title='My Tailored Resumes')


@bp.route('/submit-tailored-feedback/<job_id>', methods=['POST'])
@login_required
def submit_tailored_feedback(job_id):
    """Accept feedback (rating, comments) for a tailored resume/cover letter for a job."""
    if not current_user.is_applicant():
        return jsonify({'error': 'Unauthorized'}), 403
    mongo_db = current_app.mongo_db
    tailored = mongo_db.tailored_resumes.find_one({"user_id": str(current_user.id), "job_id": str(job_id)})
    if not tailored:
        return jsonify({'error': 'No tailored resume found for this job.'}), 404
    data = request.get_json() or {}
    rating = data.get('rating')  # e.g., 1-5
    comments = data.get('comments', '').strip()
    feedback_entry = {
        'rating': rating,
        'comments': comments,
        'submitted_at': datetime.utcnow()
    }
    # Store feedback as a list (history) in the tailored_resumes document
    mongo_db.tailored_resumes.update_one(
        {"user_id": str(current_user.id), "job_id": str(job_id)},
        {"$push": {"feedback": feedback_entry}}
    )
    return jsonify({'success': True, 'message': 'Feedback submitted.'})


@bp.route('/send-application/<job_id>', methods=['POST'])
@login_required
def send_application(job_id):
    """Send application email with resume and cover letter"""
    try:
        mongo_db = current_app.mongo_db
        job = mongo_db.jobs.find_one({"_id": ObjectId(job_id)})
        if not job:
            return jsonify({"success": False, "error": "Job not found"})
        
        # Check if job is active (not paused or closed)
        job_status = job.get('status', 'active')  # Default to active for legacy data
        if job_status not in ['active']:
            return jsonify({"success": False, "error": "This job is no longer accepting applications"})
        
        # Get application email from job data
        application_email = job.get('email') or job.get('application_email') or job.get('contact_email')
        if not application_email:
            return jsonify({"success": False, "error": "No application email found for this job"})
        
        data = request.json
        resume_text = data.get('resume_text')
        cover_letter_text = data.get('cover_letter_text')
        
        if not resume_text or not cover_letter_text:
            return jsonify({"success": False, "error": "Resume or cover letter is missing"})
        
        # Log the application attempt
        current_app.logger.info(f"User {current_user.id} is sending application for job {job_id}")
        
        # Generate PDF attachments from the text content
        resume_pdf = io.BytesIO()
        cover_letter_pdf = io.BytesIO()
        
        # Create resume PDF
        p = canvas.Canvas(resume_pdf, pagesize=letter)
        p.setTitle(f"Resume for {job.get('title')} at {job.get('company')}")
        text_object = p.beginText(72, 750)  # 1 inch from top
        text_object.setFont("Helvetica", 12)
        for line in resume_text.split('\n'):
            text_object.textLine(line)
        p.drawText(text_object)
        p.save()
        resume_pdf.seek(0)
        
        # Create cover letter PDF
        p = canvas.Canvas(cover_letter_pdf, pagesize=letter)
        p.setTitle(f"Cover Letter for {job.get('title')} at {job.get('company')}")
        text_object = p.beginText(72, 750)  # 1 inch from top
        text_object.setFont("Helvetica", 12)
        for line in cover_letter_text.split('\n'):
            text_object.textLine(line)
        p.drawText(text_object)
        p.save()
        cover_letter_pdf.seek(0)
        
        # Prepare email
        msg = MIMEMultipart()
        msg['From'] = current_user.email
        msg['To'] = application_email
        msg['Subject'] = f"Application for {job.get('title')} position"
        
        # Email body
        body = f"""Dear Hiring Manager,

I am writing to express my interest in the {job.get('title')} position at {job.get('company')}.

Please find my resume and cover letter attached to this email.

Thank you for your consideration.

Sincerely,
{current_user.first_name} {current_user.last_name}
{current_user.email}
"""
        msg.attach(MIMEText(body, 'plain'))
        
        # Attach resume
        resume_attachment = MIMEApplication(resume_pdf.read(), _subtype="pdf")
        resume_attachment.add_header('Content-Disposition', f'attachment; filename=resume_{current_user.last_name}.pdf')
        msg.attach(resume_attachment)
        
        # Attach cover letter
        cover_letter_attachment = MIMEApplication(cover_letter_pdf.read(), _subtype="pdf")
        cover_letter_attachment.add_header('Content-Disposition', f'attachment; filename=cover_letter_{current_user.last_name}.pdf')
        msg.attach(cover_letter_attachment)
        
        # Get SMTP settings from config
        smtp_server = current_app.config.get('SMTP_SERVER', 'smtp.gmail.com')
        smtp_port = current_app.config.get('SMTP_PORT', 587)
        smtp_user = current_app.config.get('SMTP_USER', current_user.email)
        smtp_password = current_app.config.get('SMTP_PASSWORD', '')
        
        # Send email through SMTP
        if current_app.config.get('TESTING') or not smtp_password:
            # In testing mode or if no SMTP password is configured, just log the attempt
            current_app.logger.info(f"Email would be sent to {application_email} for job {job.get('title')}")
        else:
            # Real email sending
            with smtplib.SMTP(smtp_server, smtp_port) as server:
                server.starttls()
                server.login(smtp_user, smtp_password)
                server.send_message(msg)
        
        # Record this application in the database
        mongo_db.applications.insert_one({
            "user_id": str(current_user.id),
            "job_id": str(job_id),
            "applied_at": datetime.utcnow(),
            "method": "email",
            "status": "applied",
            "email": application_email,
            "job_title": job.get('title'),
            "company": job.get('company')
        })
        
        return jsonify({
            "success": True, 
            "message": f"Application sent successfully to {application_email}"
        })
        
    except Exception as e:
        current_app.logger.error(f"Error sending application email: {str(e)}", exc_info=True)
        return jsonify({
            "success": False, 
            "error": "An error occurred while sending your application. Please try again."
        })


@bp.route('/talent-journey')
@login_required
def talent_journey():
    """Talent Journey (Hiring Pipeline) page"""
    candidates = []
    
    if current_user.is_recruiter():
        try:
            # Get all job postings by current recruiter
            recruiter_jobs = current_user.job_postings.all()
            job_ids = [job.id for job in recruiter_jobs]
            
            # Get applications to recruiter's jobs if any exist
            if job_ids:
                applications = Application.query.filter(
                    Application.job_posting_id.in_(job_ids)
                ).join(User).all()
                
                # Format data for template
                for app in applications:
                    # Get ATS score from match_score field
                    ats_score = app.match_score if app.match_score is not None else None
                    
                    # Format ATS score display
                    if ats_score is not None:
                        ats_score_display = f"{int(ats_score)}%"
                        if ats_score >= 80:
                            ats_score_class = "text-success fw-bold"
                        elif ats_score >= 60:
                            ats_score_class = "text-warning fw-bold"
                        else:
                            ats_score_class = "text-danger fw-bold"
                        print(f"   Final display: {ats_score_display} (class: {ats_score_class})")
                    else:
                        ats_score_display = "Not Available"
                        ats_score_class = "text-muted"
                        print(f"   Final display: {ats_score_display} (class: {ats_score_class})")
                    
                    candidates.append({
                        'application_id': app.id,
                        'name': f"{app.user.first_name or ''} {app.user.last_name or ''}".strip() or app.user.email,
                        'job': getattr(app, 'job_title', 'N/A'),
                        'status': app.status,
                        'applied_date': app.created_at.strftime('%b %d, %Y') if app.created_at else 'N/A',
                        'ats_score': ats_score,
                        'ats_score_display': ats_score_display,
                        'ats_score_class': ats_score_class
                    })
                    
                    # DEBUG: Print what we're adding
                    print(f"   Added candidate: {app.user.email} - Status: '{app.status}' - Job: {app.job_title}")
        except Exception as e:
            print(f"Error in talent_journey: {e}")
    
    return render_template('talent_journey.html', candidates=candidates)


@bp.route('/api/update-application-status', methods=['POST'])
@login_required
def update_application_status():
    """API endpoint to update application status"""
    try:
        data = request.get_json()
        application_id = data.get('application_id')
        new_status = data.get('status')
        
        if not application_id or not new_status:
            return jsonify({'success': False, 'error': 'Missing application_id or status'})
        
        # Get the application
        application = Application.query.get(application_id)
        if not application:
            return jsonify({'success': False, 'error': 'Application not found'})
        
        # Check if current user is the recruiter for this job
        # Skip authorization check for now - recruiters can update any application status
        if not current_user.is_recruiter():
            return jsonify({'success': False, 'error': 'Unauthorized - Only recruiters can update status'})
        
        # Update the status
        application.status = new_status
        db.session.commit()
        
        return jsonify({'success': True, 'message': 'Status updated successfully'})
        
    except Exception as e:
        current_app.logger.error(f"Error updating application status: {str(e)}")
        return jsonify({'success': False, 'error': 'Internal server error'})


@bp.route('/api/update-my-application-status', methods=['POST'])
@login_required
def update_my_application_status():
    """API endpoint for applicants to update their own application status"""
    try:
        current_app.logger.info(f"Update application status request from user {current_user.id}")
        
        data = request.get_json()
        if not data:
            current_app.logger.error("No JSON data received")
            return jsonify({'success': False, 'error': 'No data received'})
        
        application_id = data.get('application_id')
        new_status = data.get('status')
        source = data.get('source', 'postgres')  # 'postgres' or 'mongodb'
        notes = data.get('notes', '')
        
        current_app.logger.info(f"Request data: application_id={application_id}, status={new_status}, source={source}")
        
        if not application_id or not new_status:
            current_app.logger.error(f"Missing required fields: application_id={application_id}, status={new_status}")
            return jsonify({'success': False, 'error': 'Missing application_id or status'})
        
        # Handle prefixed IDs from combined applications list
        actual_app_id = application_id
        if isinstance(application_id, str):
            if application_id.startswith('pg_'):
                actual_app_id = application_id[3:]  # Remove 'pg_' prefix
                source = 'postgres'
            elif application_id.startswith('mg_'):
                actual_app_id = application_id[3:]  # Remove 'mg_' prefix
                source = 'mongodb'
            elif len(application_id) == 24:
                # Looks like a MongoDB ObjectId
                source = 'mongodb'
            else:
                # Try to convert to int to see if it's a PostgreSQL ID
                try:
                    int(application_id)
                    source = 'postgres'
                except ValueError:
                    # If it's not an integer, assume it's MongoDB
                    source = 'mongodb'
        
        current_app.logger.info(f"Processed application_id: {actual_app_id}, determined source: {source}")
        
        # Validate status
        valid_statuses = ['applied', 'screening', 'interview_scheduled', 'interviewed', 
                         'offer_received', 'accepted', 'rejected', 'withdrawn', 'no_response']
        if new_status not in valid_statuses:
            current_app.logger.error(f"Invalid status: {new_status}")
            return jsonify({'success': False, 'error': 'Invalid status'})
        
        if source == 'postgres':
            # Update PostgreSQL application
            from app.models.application import Application
            current_app.logger.info(f"Looking up PostgreSQL application with ID: {actual_app_id}")
            
            application = Application.query.get(actual_app_id)
            if not application:
                current_app.logger.error(f"Application not found: {actual_app_id}")
                return jsonify({'success': False, 'error': 'Application not found'})
            
            if application.user_id != current_user.id:
                current_app.logger.error(f"Unauthorized access: application user_id={application.user_id}, current_user_id={current_user.id}")
                return jsonify({'success': False, 'error': 'Unauthorized access to application'})
            
            old_status = application.status
            current_app.logger.info(f"Updating application status from {old_status} to {new_status}")
            
            application.update_status(new_status, notes)
            current_app.logger.info(f"Successfully updated application status")
            
            return jsonify({
                'success': True, 
                'message': f'Application status updated from {old_status} to {new_status}',
                'new_status': new_status
            })
            
        else:  # MongoDB application
            try:
                mongo_db = get_mongo_db()
                current_app.logger.info(f"MongoDB connection obtained")
            except Exception as e:
                current_app.logger.error(f"MongoDB connection failed: {str(e)}")
                return jsonify({'success': False, 'error': 'Database connection error'})
            
            current_app.logger.info(f"Looking up MongoDB application with ID: {actual_app_id}")
            
            # Update MongoDB application
            result = mongo_db.job_applications.update_one(
                {
                    '_id': ObjectId(actual_app_id),
                    'user_id': str(current_user.id)
                },
                {
                    '$set': {
                        'status': new_status,
                        'updated_at': datetime.utcnow()
                    }
                }
            )
            
            if result.matched_count == 0:
                current_app.logger.error(f"MongoDB application not found or unauthorized: {actual_app_id}")
                return jsonify({'success': False, 'error': 'Application not found or unauthorized'})
            
            current_app.logger.info(f"Successfully updated MongoDB application status to {new_status}")
            
            return jsonify({
                'success': True, 
                'message': f'Application status updated to {new_status}',
                'new_status': new_status
            })
        
    except Exception as e:
        current_app.logger.error(f"Error updating my application status: {str(e)}")
        return jsonify({'success': False, 'error': 'Internal server error'})


# --- Interview Questions Integration ---
@bp.route('/tailor/<job_id>/database-questions')
@login_required
def tailor_database_questions(job_id):
    """Generate interview questions for specific job"""
    try:
        # Get MongoDB database from current app
        mongo_db = current_app.mongo_db
        
        if mongo_db is None:
            flash('Database connection not available', 'error')
            return redirect(url_for('main.home'))
        
        # Find the job
        jobs_collection = mongo_db.jobs
        job = jobs_collection.find_one({"_id": ObjectId(job_id)})
        
        if job is None:
            flash('Job not found', 'error')
            return redirect(url_for('main.home'))
        
        # Check if this is for iframe usage
        is_iframe = request.args.get('iframe') == 'true'
        template_name = 'question/tailor_database_questions_iframe.html' if is_iframe else 'question/tailor_database_questions.html'
            
        return render_template(template_name, job=job, job_id=job_id)
    except Exception as e:
        current_app.logger.error(f'Error loading questions page: {str(e)}')
        flash(f'Error loading questions page: {str(e)}', 'error')
        return redirect(url_for('main.tailor_resume', job_id=job_id))


@bp.route('/tailor/<job_id>/generate-database-questions', methods=['POST'])
@login_required
def generate_database_questions_route(job_id):
    """Generate questions endpoint"""
    try:
        # Import the question generation functionality
        from app.question.question_gen_db import generate_database_questions
        
        # Get form data
        num_questions = int(request.form.get('num_questions', 5))
        question_type = request.form.get('question_type', 'mixed')
        
        # Get MongoDB database from current app
        mongo_db = current_app.mongo_db
        
        if mongo_db is None:
            flash('Database connection not available', 'error')
            return redirect(url_for('main.tailor_database_questions', job_id=job_id))
        
        # Find the job
        jobs_collection = mongo_db.jobs
        job = jobs_collection.find_one({"_id": ObjectId(job_id)})
        
        if job is None:
            flash('Job not found', 'error')
            return redirect(url_for('main.tailor_database_questions', job_id=job_id))
            
        # Generate questions with parameters
        raw_questions_data = generate_database_questions(
            job_id=job_id, 
            job_data=job, 
            n=num_questions,
            question_type=question_type
        )
        
        # Process questions data using utility function for better parsing
        questions_data = build_questions_data(raw_questions_data) if raw_questions_data else []
        
        # Create question statistics for display
        question_stats = None
        if questions_data and isinstance(questions_data, list):
            code_snippets = sum(1 for q in questions_data if q.get('code_snippet'))
            total_words = sum(len(str(q.get('text', '') + q.get('expected', '') + q.get('relevance', '')).split()) for q in questions_data)
            question_stats = {
                'question_count': len(questions_data),
                'code_snippets': code_snippets,
                'word_count': total_words,
                'estimated_reading_time': max(1, total_words // 200)  # ~200 words per minute
            }
        
        # Check if this is for iframe usage
        is_iframe = request.args.get('iframe') == 'true'
        template_name = 'question/tailor_database_questions_iframe.html' if is_iframe else 'question/tailor_database_questions.html'
        
        return render_template(template_name, 
                             job=job, 
                             job_id=job_id, 
                             questions=questions_data,
                             num_questions=len(questions_data) if questions_data else 0,
                             question_type=question_type,
                             question_stats=question_stats)
        
    except Exception as e:
        current_app.logger.error(f'Error generating questions: {str(e)}')
        flash(f'Error generating questions: {str(e)}', 'error')
        return redirect(url_for('main.tailor_database_questions', job_id=job_id))


# --- Interview Questions Main Routes ---
@bp.route('/interview-questions')
@login_required
def interview_questions_home():
    """Main Interview Questions page"""
    return render_template('question/index.html')


@bp.route('/interview-questions/job-skills')
# @login_required  # TEMPORARY: Removed for testing
def interview_questions_skills():
    """Interview Questions based on job skills"""
    return render_template('question/skills_questions.html')


@bp.route('/interview-questions/job-description')
@login_required
def interview_questions_description():
    """Interview Questions based on job description"""
    return render_template('question/job_description_questions.html')


@bp.route('/interview-questions/from-database')
@login_required
def interview_questions_database():
    """Interview Questions from database"""
    return render_template('question/questions_from_db.html')


@bp.route('/api/generate-questions-skills', methods=['POST'])
# @login_required  # TEMPORARY: Removed for testing
def generate_questions_from_skills():
    """Generate questions based on skills"""
    try:
        from app.question.question_gen import generate_questions_from_skills as gen_skills
        
        # Handle both JSON and form data
        if request.is_json:
            data = request.get_json()
        else:
            data = request.form.to_dict()
            
        skills = data.get('skills', '')
        level = data.get('level', 'intermediate')
        question_type = data.get('question_type', 'technical')
        language = data.get('language', 'English')
        num_questions = int(data.get('num_questions', data.get('count', 5)))
        mode = data.get('mode', '')  # Check if this is recruiter mode
        if not skills:
            if request.is_json:
                return jsonify({'error': 'Skills are required', 'success': False}), 400
            else:
                flash('Skills are required', 'error')
                # Redirect based on mode
                if mode == 'recruiter':
                    return redirect(url_for('recruiter.recruiter_interview'))
                else:
                    return redirect(url_for('main.interview_questions_skills'))
                
        # Generate questions and get structured data
        raw_questions_data = gen_skills(skills, level, question_type, language, num_questions)

        # Already returns list[dict]; build_questions_data expects uniform keys; keep as-is if structured
        if raw_questions_data and isinstance(raw_questions_data, list) and raw_questions_data[0] and 'text' in raw_questions_data[0]:
            questions_data = raw_questions_data
        else:
            questions_data = build_questions_data(raw_questions_data) if raw_questions_data else []

        # Create question statistics for display
        question_stats = None
        if questions_data and isinstance(questions_data, list):
            code_snippets = sum(1 for q in questions_data if q.get('code_snippet'))
            total_words = sum(len(str(q.get('text', '') + q.get('expected', '') + q.get('relevance', '')).split()) for q in questions_data)
            question_stats = {
                'question_count': len(questions_data),
                'code_snippets': code_snippets,
                'word_count': total_words,
                'estimated_reading_time': max(1, total_words // 200)  # ~200 words per minute
            }

        actual_count = len(questions_data)
        placeholder_count = sum(1 for q in questions_data if q.get('text','').lower().startswith('placeholder'))
        # Normalizar tipo (conservar original si viene capitalizado del form en otras vistas)
        qt_display = question_type
        if isinstance(qt_display, str):
            qt_display = qt_display.strip()
        if request.is_json:
            return jsonify({
                'questions': questions_data,
                'requested_num_questions': num_questions,
                'returned_num_questions': actual_count,
                'placeholder_count': placeholder_count,
                'question_type': qt_display,
                'success': True
            })
        else:
            return render_template('question/skills_questions.html', 
                                   questions=questions_data, 
                                   form_data={**data, 'num_questions': actual_count, 'question_type': qt_display},
                                   num_questions=actual_count,
                                   requested_num_questions=num_questions,
                                   returned_num_questions=actual_count,
                                   placeholder_count=placeholder_count,
                                   question_type=qt_display,
                                   question_stats=question_stats,
                                   mode=mode)  # Pass mode to template
        
    except Exception as e:
        current_app.logger.error(f'Error generating skills questions: {str(e)}')
        if request.is_json:
            return jsonify({'error': str(e), 'success': False}), 500
        else:
            flash(f'Error generating questions: {str(e)}', 'error')
            # Redirect based on mode
            mode = request.form.get('mode', '') if request.form else ''
            if mode == 'recruiter':
                return redirect(url_for('recruiter.recruiter_interview'))
            else:
                return redirect(url_for('main.interview_questions_skills'))


@bp.route('/api/generate-questions-description', methods=['POST'])
@login_required
def generate_questions_from_description():
    """Generate questions based on job description"""
    try:
        from app.question.question_gen2 import question_generator_for_ui as gen_desc
        
        # Handle both JSON and form data
        if request.is_json:
            data = request.get_json()
        else:
            data = request.form.to_dict()
            
        job_description = data.get('job_description', '')
        role = data.get('job_position', 'Developer')  # Form uses job_position
        level = data.get('level', 'intermediate')
        previous_experience = data.get('previous_experience', f'{level} level professional')
        question_type = data.get('question_type', 'technical')
        language = data.get('language', 'English')
        num_questions = int(data.get('num_questions', data.get('count', 5)))
        
        if not job_description:
            if request.is_json:
                return jsonify({'error': 'Job description is required', 'success': False}), 400
            else:
                flash('Job description is required', 'error')
                return redirect(url_for('main.interview_questions_description'))
                
        # Generate questions and get structured data
        raw_questions_data = gen_desc(job_description, role, level, previous_experience, question_type, language, num_questions)

        # Already returns list[dict]; build_questions_data expects uniform keys; keep as-is if structured
        if raw_questions_data and isinstance(raw_questions_data, list) and 'text' in raw_questions_data[0]:
            questions_data = raw_questions_data
        else:
            questions_data = build_questions_data(raw_questions_data) if raw_questions_data else []
        
        # Create question statistics for display
        question_stats = None
        if questions_data and isinstance(questions_data, list):
            code_snippets = sum(1 for q in questions_data if q.get('code_snippet'))
            total_words = sum(len(str(q.get('text', '') + q.get('expected', '') + q.get('relevance', '')).split()) for q in questions_data)
            question_stats = {
                'question_count': len(questions_data),
                'code_snippets': code_snippets,
                'word_count': total_words,
                'estimated_reading_time': max(1, total_words // 200)  # ~200 words per minute
            }
        # Counts & placeholder meta (mirrors skills route)
        actual_count = len(questions_data)
        placeholder_count = sum(1 for q in questions_data if q.get('text','').lower().startswith('placeholder'))

        # Normalize display of question type
        qt_display = question_type
        if isinstance(qt_display, str):
            qt_display = qt_display.strip()

        if request.is_json:
            return jsonify({
                'questions': questions_data,
                'requested_num_questions': num_questions,
                'returned_num_questions': actual_count,
                'placeholder_count': placeholder_count,
                'question_type': qt_display,
                'success': True
            })
        else:
            return render_template('question/job_description_questions.html', 
                                   questions=questions_data, 
                                   form_data={**data, 'num_questions': actual_count, 'question_type': qt_display},
                                   num_questions=actual_count,
                                   requested_num_questions=num_questions,
                                   returned_num_questions=actual_count,
                                   placeholder_count=placeholder_count,
                                   question_type=qt_display,
                                   question_stats=question_stats)
        
    except Exception as e:
        current_app.logger.error(f'Error generating description questions: {str(e)}')
        if request.is_json:
            return jsonify({'error': str(e), 'success': False}), 500
        else:
            flash(f'Error generating questions: {str(e)}', 'error')
            return redirect(url_for('main.interview_questions_description'))
