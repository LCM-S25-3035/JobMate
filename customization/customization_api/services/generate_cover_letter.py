import os
import json
from docx import Document
from customization.customization_api.llm_clients.gemini import generate_from_gemini
from customization.customization_api.utils.file_loader import load_json, load_template_text

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../../"))
TEMPLATE_PATH = os.path.join(BASE, "customization/template/cover_letter_template1.docx")
OUTPUT_PATH = os.path.join(BASE, "cover_letter/generated_cover_letter.docx")

def generate_prompt(resume, job, template_text):
    name = resume["personal_information"]["name"]
    email = resume["personal_information"]["email"]
    education = ", ".join([edu["degree"] for edu in resume["education"]])
    matched_skills = ", ".join(resume["soft_skills"] + resume["technical_skills"])
    job_title = job["job_title"]
    company = job["company"]["name"]

    return f"""
    Below is a resume and a job description. Please generate a professional, ATS-friendly cover letter by filling in the missing parts of the following template. Replace placeholders like [job title], [organisation], [your name], [your email address], etc.

    === Resume Summary ===
    Name: {name}
    Email: {email}
    Education: {education}
    Experience: {resume["professional_summary"]}
    Skills: {matched_skills}

    === Job Description ===
    Role: {job_title}
    Company: {company}
    Description: {job["job_description"]}
    Requirements: {", ".join(job["requirements"])}
    Responsibilities: {", ".join(job["responsibilities"])}

    === Cover Letter Template ===
    {template_text}

    Return a completed, professional version of the cover letter.
    """

def generate_cover_letter():
    resume_path = os.path.join(BASE, "resume/resume.json")
    job_path = os.path.join(BASE, "resume/job_posting.json")
    resume = load_json(resume_path)
    job = load_json(job_path)
    template_text = load_template_text(TEMPLATE_PATH)

    prompt = generate_prompt(resume, job, template_text)
    response_text = generate_from_gemini(prompt)

    # Save result to .docx
    doc = Document()
    for line in response_text.split("\n"):
        doc.add_paragraph(line)

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    doc.save(OUTPUT_PATH)

    return OUTPUT_PATH


# Reference:
# OpenAI 4o, 1st prompt:
# Given that I have all these json files, and job posting json file, how can i pass those to Gemini to customize Targeted Cover Letter for me. Give me an example code. 

# OpenAI 4o, last prompt:
# 